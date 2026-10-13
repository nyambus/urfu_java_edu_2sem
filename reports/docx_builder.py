"""Генератор отчётов по лабораторным работам в формате примера
«Отчет_о_лабораторной_работе_13_Козионов.docx».

Формат не «похож», а совпадает: скрипт берёт файл примера как шаблон,
удаляет из него содержимое и собирает новый отчёт на его же стилях
(SectionTitle, BodyText, CaptionText), его параметрах страницы и его
колонтитуле с номером страницы. Проверено рендерингом в LibreOffice:
стр. 1 — титульный лист без номера, стр. 2 — «Екатеринбург / 2026 год /
ЦЕЛЬ РАБОТЫ», далее разделы.

Зависимость: python-docx >= 1.1   ->   pip install python-docx

Демонстрация:  python docx_builder.py demo.docx
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE / "Отчет_о_лабораторной_работе_13_Козионов.docx"

FONT = "Times New Roman"
BODY_SIZE = Pt(14)       # основной текст — 14 пт (в примере sz=28)
CAPTION_SIZE = Pt(12)    # подписи рисунков — 12 пт курсив (sz=24)
TITLE_SIZE = Pt(14)      # титульный лист — 14 пт, интервал одинарный
LINE_1_5 = 1.5

TITLE_UNIVERSITY = (
    "Федеральное государственное автономное образовательное учреждение "
    "высшего образования «Уральский федеральный университет имени первого "
    "Президента России Б. Н. Ельцина»"
)
TITLE_FACULTY = "Факультет радиоэлектроники и информационных технологий — РТФ"
TITLE_DISCIPLINE = "Дисциплина «Программирование»"
TITLE_GROUP = "студент группы РИЗ-150110у"
TITLE_STUDENT = "Козионов Илья Алексеевич"
TITLE_TEACHER = "Архипов Николай Александрович"
TITLE_CITY = "Екатеринбург"
TITLE_YEAR = "2026 год"

# id стилей внутри шаблона-примера
ST_SECTION = "SectionTitle"
ST_BODY = "BodyText"
ST_CAPTION = "CaptionText"
ST_HYPERLINK = "aff8"  # <w:style w:type="character" w:styleId="aff8"> (Hyperlink)


def _hyperlink(paragraph, url: str, style_id: str = ST_HYPERLINK):
    """Настоящая гиперссылка (w:hyperlink + внешняя связь), как в примере."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), style_id)
    rpr.append(rstyle)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    return paragraph


class Report:
    """Отчёт: титульный лист, разделы, рисунки с подписями, листинги, вывод."""

    def __init__(
        self,
        lab_number: int,
        lab_title: str,
        template: str | Path | None = DEFAULT_TEMPLATE,
        student_group: str = TITLE_GROUP,
        student_name: str = TITLE_STUDENT,
        teacher: str = TITLE_TEACHER,
        city: str = TITLE_CITY,
        year: str = TITLE_YEAR,
    ):
        tpl = Path(template) if template else None
        if tpl is not None and tpl.exists():
            self.doc = Document(str(tpl))
            self._strip_body()
            self.from_template = True
        else:
            self.doc = Document()
            self._setup_page_fallback()
            self._setup_styles_fallback()
            self.from_template = False
        self.figure_no = 0
        self._title_page(lab_number, lab_title, student_group, student_name, teacher)
        self._page_break()
        # как в примере: «Екатеринбург» и «2026 год» открывают вторую страницу
        self.centered(city, spacing=1.0)
        self.centered(year, spacing=1.0)

    # --- подготовка документа ----------------------------------------------

    def _strip_body(self):
        """Удаляет содержимое примера, сохраняя sectPr, стили и колонтитул."""
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        # картинки и внешние ссылки примера больше не нужны: иначе они попадут
        # в готовый файл (лишние ~1 МБ мусора в каждом отчёте)
        for r_id, rel in list(self.doc.part.rels.items()):
            if rel.reltype.endswith("/image") or rel.is_external:
                self.doc.part.drop_rel(r_id)

    def _setup_page_fallback(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)
        section.different_first_page_header_footer = True
        self._page_number_footer(section.footer)

    @staticmethod
    def _page_number_footer(footer):
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)

    def _setup_styles_fallback(self):
        """Только если файла-примера нет рядом со скриптом."""
        doc = self.doc
        normal = doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = BODY_SIZE
        normal.paragraph_format.line_spacing = LINE_1_5

        def add(name, style_id, size, bold=False, italic=False, align=None):
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.element.set(qn("w:styleId"), style_id)
            st.element.set(qn("w:customStyle"), "1")
            st.base_style = normal
            st.font.name = FONT
            st.font.size = size
            st.font.bold = bold
            st.font.italic = italic
            if align is not None:
                st.paragraph_format.alignment = align
            st.paragraph_format.line_spacing = LINE_1_5
            st.paragraph_format.space_after = Pt(0)
            return st

        add("SectionTitle", ST_SECTION, TITLE_SIZE, bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER)
        add("BodyTextLab", ST_BODY, BODY_SIZE, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add("CaptionText", ST_CAPTION, CAPTION_SIZE, italic=True,
            align=WD_ALIGN_PARAGRAPH.CENTER)
        link = doc.styles.add_style("HyperLink", WD_STYLE_TYPE.CHARACTER)
        link.element.set(qn("w:styleId"), ST_HYPERLINK)
        link.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
        link.font.underline = True

    # --- титульный лист ----------------------------------------------------

    def _title_page(self, lab_number, lab_title, group, name, teacher):
        self.centered(TITLE_UNIVERSITY, spacing=1.0)
        for _ in range(4):
            self.centered("", spacing=1.0)
        self.centered(TITLE_FACULTY, spacing=1.0)
        self.centered(TITLE_DISCIPLINE, spacing=1.0)
        for _ in range(2):
            self.centered("", spacing=1.0)
        self.centered(f"Отчет к лабораторной работе № {lab_number}", spacing=1.0)
        self.centered(f"«{lab_title}»", spacing=1.0)
        for _ in range(6):
            self.centered("", spacing=1.0)
        indent = Emu(5245 * 635)  # 5245 twips = абзацный отступ блока «Выполнил»
        self.centered("Выполнил:", spacing=1.0, indent=indent)
        self.centered(group, spacing=1.0, indent=indent)
        self.centered(name, spacing=1.0, indent=indent)
        self.centered("", spacing=1.0, indent=indent)
        self.centered("Принял:", spacing=1.0, indent=indent)
        self.centered("Преподаватель", spacing=1.0, indent=indent)
        self.centered(teacher, spacing=1.0, indent=indent)

    # --- примитивы ---------------------------------------------------------

    def _page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(10)
        p.add_run().add_break(WD_BREAK.PAGE)

    def centered(self, text, spacing=LINE_1_5, indent=None, bold=False):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = spacing
        p.paragraph_format.space_after = Pt(0)
        if indent is not None:
            p.paragraph_format.first_line_indent = indent
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = TITLE_SIZE if spacing == 1.0 else BODY_SIZE
        run.font.bold = bold
        return p

    def section(self, text):
        """Заголовок раздела: ЦЕЛЬ РАБОТЫ / ЗАДАНИЕ №N / ВЫВОД."""
        p = self.doc.add_paragraph(text, style=ST_SECTION)
        self.doc.add_paragraph("", style=ST_SECTION)
        return p

    def body(self, text):
        """Абзац основного текста: 14 пт, по ширине, без отступа, 1,5."""
        p = self.doc.add_paragraph(text, style=ST_BODY)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Pt(0)
        for r in p.runs:
            r.font.name = FONT
            r.font.size = BODY_SIZE
        return p

    def gap(self):
        return self.doc.add_paragraph("", style=ST_BODY)

    def link(self, url):
        """Ссылка на файл в репозитории — отдельным абзацем по центру."""
        p = self.doc.add_paragraph(style=ST_BODY)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _hyperlink(p, url)
        return p

    def image(self, path=None, caption="", width_in=6.0):
        """Рисунок + подпись «Рисунок N — <caption>».

        path=None -> вставляется заметный заполнитель, чтобы отчёт можно было
        собрать до появления скриншотов.
        """
        self.figure_no += 1
        p = self.doc.add_paragraph(style=ST_BODY)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path and Path(path).exists():
            p.add_run().add_picture(str(path), width=Inches(width_in))
        else:
            run = p.add_run(f"[СКРИНШОТ: {caption}]")
            run.font.name = FONT
            run.font.size = BODY_SIZE
            run.font.italic = True
        return self.doc.add_paragraph(
            f"Рисунок {self.figure_no} — {caption}", style=ST_CAPTION
        )

    def code(self, text, caption=None, size=Pt(11)):
        """Листинг кода (альтернатива скриншоту, допустима по требованиям ЛР)."""
        if caption:
            self.doc.add_paragraph(caption, style=ST_BODY)
        for line in text.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph(style=ST_BODY)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line if line.strip() else " ")
            run.font.size = size
            rpr = run._r.get_or_add_rPr()
            rfonts = OxmlElement("w:rFonts")
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rfonts.set(qn(attr), "Consolas")
            rpr.append(rfonts)

    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return str(path)


def demo(path="demo.docx"):
    rep = Report(2, "Разработка простого REST сервиса")
    rep.section("ЦЕЛЬ РАБОТЫ")
    rep.body("Цель работы: разработать простой REST сервис.")
    rep.gap()
    rep.section("ЗАДАНИЕ №1")
    rep.body("Реализовать простой REST сервис с валидацией.")
    rep.link("https://github.com/USER/urfu-spring-labs/blob/main/lab2/pom.xml")
    rep.image(None, "Класс MyController")
    rep.code("public class MyController {\n    // ...\n}", caption="Листинг 1 — MyController")
    rep.gap()
    rep.section("ВЫВОД")
    rep.body("Вывод: сервис реализован и протестирован.")
    return rep.save(path)


if __name__ == "__main__":
    print(demo(sys.argv[1] if len(sys.argv) > 1 else "demo.docx"))
